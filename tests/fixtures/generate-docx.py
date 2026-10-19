"""Generate synthetic interoperability fixtures using python-docx 1.2.0, not our ZIP writer."""
from pathlib import Path
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

out = Path(__file__).parent / "office"
out.mkdir(exist_ok=True)
doc = Document()
doc.core_properties.author = "PrivateAI synthetic fixture"
doc.core_properties.title = "Synthetic equipment handover"
doc.add_heading("Synthetic equipment handover", 0)
p = doc.add_paragraph()
p.add_run("Acknowledgement must arrive ")
p.add_run("strictly before 09:45 UTC").bold = True
p.add_run(" on 21 October 2026. Arrival at 09:45 does not qualify.")
doc.add_paragraph("Keep the guard secured during use; disconnect power before maintenance; report damage to the supervisor.")
table = doc.add_table(rows=1, cols=3)
for cell, value in zip(table.rows[0].cells, ["Item", "Quantity", "Unit price"]): cell.text = value
for values in [("Adapter", "3", "EUR 14.20"), ("Cable", "2", "EUR 5.75")]:
    for cell, value in zip(table.add_row().cells, values): cell.text = value
doc.add_paragraph("These prices exclude tax. Tax rate and delivery cost are not supplied.")
doc.add_paragraph("Check the serial number.", style="List Number")
doc.add_paragraph("Record the handover time.", style="List Number")
doc.sections[0].header.paragraphs[0].text = "SYNTHETIC_HEADER_NOT_IN_BODY"
doc.sections[0].footer.paragraphs[0].text = "SYNTHETIC_FOOTER_NOT_IN_BODY"
doc.save(out / "handover-python-docx.docx")
# An independently packaged revision should fail rather than combine old/new text.
revision = Document()
p = revision.add_paragraph("Unchanged context. ")
insert = OxmlElement("w:ins")
insert.set(qn("w:id"), "1")
r = OxmlElement("w:r")
t = OxmlElement("w:t")
t.text = "UNRESOLVED_REVISION"
r.append(t); insert.append(r); p._p.append(insert)
revision.save(out / "revision-python-docx.docx")
